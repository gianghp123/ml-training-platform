import json
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / "Tai_lieu_block_pipeline_ML.docx"
blocks = json.loads((ROOT / "blocks.json").read_text(encoding="utf-8"))

CATEGORY_ORDER = ["load_data", "preprocess_data", "split_data", "transform", "config", "model", "evaluate", "export"]
CATEGORY_INFO = {
    "load_data": ("1", "Load Data", "Nguồn dữ liệu đầu vào của pipeline."),
    "preprocess_data": ("2", "Preprocess", "Làm sạch và chuẩn hóa dữ liệu."),
    "split_data": ("3", "Split Data", "Chia dữ liệu thành train/test hoặc folds."),
    "transform": ("4", "Transform", "Biến đổi và tạo đặc trưng."),
    "config": ("5", "Configuration", "Cung cấp cấu hình huấn luyện qua các artifact riêng."),
    "model": ("6", "Model", "Huấn luyện mô hình phân loại, hồi quy, phân cụm và deep learning."),
    "evaluate": ("7", "Evaluate", "Đánh giá mô hình và tạo kết quả dự đoán."),
    "export": ("8", "Export", "Lưu model, dataset hoặc metrics ra hệ thống ngoài."),
}
ARTIFACTS = {
    "Dataset": "Dữ liệu dạng bảng dùng cho tiền xử lý, chia tập, huấn luyện và đánh giá.",
    "Folds": "Tập mô tả các fold phục vụ cross-validation.",
    "TrainedModel": "Mô hình đã hoàn thành huấn luyện và có thể dự đoán/lưu registry.",
    "ModelSpec": "Đặc tả thuật toán và siêu tham số để tạo lại estimator cho từng fold.",
    "Predictions": "Bảng kết quả dự đoán, xác suất và sai số theo từng bản ghi.",
    "FittedTransformer": "Bộ biến đổi đã fit, dùng để áp dụng cùng phép biến đổi cho validation/test.",
    "Hyperparameters": "Cấu hình learning rate, epochs, batch size, optimizer và weight decay.",
    "LossConfig": "Cấu hình hàm mất mát theo loại bài toán.",
    "EarlyStoppingConfig": "Cấu hình điều kiện dừng sớm trong huấn luyện.",
    "Metrics": "Các chỉ số và báo cáo đánh giá mô hình.",
}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "687386"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
        prevent_row_split(row)


def set_run(run, size=10.5, bold=False, color="000000", font="Calibri", italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        paragraph.add_run()._r.append(node)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, color, before, after in (
        ("Title", 28, DARK_BLUE, 0, 8),
        ("Subtitle", 14, MID_GRAY, 0, 20),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = "ML TRAINING PLATFORM  |  BLOCK & SCHEMA REFERENCE"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        set_run(run, size=8, bold=True, color=MID_GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.left_indent = Inches(0)
    footer.paragraph_format.right_indent = Inches(0)
    footer.paragraph_format.first_line_indent = Inches(0)
    # Leading non-breaking spaces protect the first visible footer glyphs in
    # Word's even-page PDF export while remaining visually neutral.
    run = footer.add_run("\u00a0\u00a0Trang ")
    set_run(run, size=8.5, color=MID_GRAY)
    add_field(footer, "PAGE")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def format_table(table, header=True, font_size=8.5):
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if header and r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run(run, size=font_size, bold=(header and r_idx == 0), color=DARK_BLUE if header and r_idx == 0 else "000000")
    if header:
        set_repeat_table_header(table.rows[0])


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + ": ")
    set_run(r, bold=True, color=DARK_BLUE)
    r = p.add_run(str(value))
    set_run(r, font="Consolas" if label in ("ID", "Code", "Source") else "Calibri")


def format_default(value):
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, list):
        return ", ".join(map(str, value))
    return str(value)


def validation_text(field):
    rules = field.get("validation", {}) or {}
    parts = []
    if rules.get("required"):
        parts.append("bắt buộc")
    for key, label in (("min", "min"), ("max", "max"), ("minLength", "độ dài min"), ("maxLength", "độ dài max"), ("step", "step")):
        if key in rules:
            parts.append(f"{label}: {rules[key]}")
    options = field.get("options") or []
    if options:
        parts.append("options: " + ", ".join(str(o.get("value")) for o in options))
    depends = field.get("dependsOn")
    if depends:
        parts.append(f"hiện khi {depends.get('field')} = {depends.get('equals')}")
    return "; ".join(parts) if parts else "—"


def add_ports_table(doc, ports):
    doc.add_paragraph("Input/Output schema", style="Heading 3")
    if not ports:
        add_bullet(doc, "Block không khai báo port.")
        return
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Port ID", "Nhãn", "Hướng", "Artifact", "Bắt buộc", "Kết nối"]
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    for port in ports:
        cells = table.add_row().cells
        values = [
            port.get("id", ""), port.get("label", ""),
            "Input" if port.get("direction") == "input" else "Output",
            port.get("artifact", ""),
            "Có" if port.get("required") else "Không",
            "Nhiều" if port.get("multiple") else "Một",
        ]
        for i, value in enumerate(values):
            cells[i].text = str(value)
    set_table_geometry(table, [1440, 2160, 1080, 1800, 1080, 1800])
    format_table(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_config_table(doc, schema):
    doc.add_paragraph("Config schema", style="Heading 3")
    if not schema:
        add_bullet(doc, "Block không có tham số cấu hình.")
        return
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, value in enumerate(["Khóa / kiểu", "Nhãn", "Mặc định", "Validation và options"]):
        table.rows[0].cells[i].text = value
    for key, field in schema.items():
        cells = table.add_row().cells
        values = [f"{key}\n[{field.get('type', 'unknown')}]", field.get("label", ""), format_default(field.get("default")), validation_text(field)]
        for i, value in enumerate(values):
            cells[i].text = str(value)
    set_table_geometry(table, [1800, 2160, 1800, 3600])
    format_table(table)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
configure_styles(doc)
add_header_footer(section)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(72)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TÀI LIỆU THIẾT KẾ")
set_run(r, size=11, bold=True, color=BLUE)
p = doc.add_paragraph("HỆ THỐNG BLOCK PIPELINE ML", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Mô tả chi tiết block, config schema và input/output schema", style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph().paragraph_format.space_after = Pt(86)
for label, value in (("Dự án", "ML Training Platform"), ("Phiên bản tài liệu", "1.0"), ("Nguồn", "Frontend block definitions"), ("Ngày cập nhật", "11/07/2026")):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + ": ")
    set_run(r, size=10.5, bold=True, color=DARK_BLUE)
    r = p.add_run(value)
    set_run(r, size=10.5, color=MID_GRAY)

doc.add_page_break()
doc.add_heading("Giới thiệu và phạm vi", level=1)
doc.add_paragraph(
    "Tài liệu này mô tả catalog block hiện có trong frontend Pipeline Builder. Nội dung được sinh trực tiếp từ các file definition tại thời điểm phát hành tài liệu, nhằm tạo một hợp đồng chung giữa giao diện, API, Pipeline Engine và runtime của từng block."
)
add_bullet(doc, "Port input chỉ nhận edge có artifact trùng với artifact của port output nguồn.")
add_bullet(doc, "Input có required = true phải được nối trước khi Save/Run.")
add_bullet(doc, "multiple = false giới hạn một edge vào port; multiple = true cho phép fan-out hoặc nhiều kết nối theo hướng khai báo.")
add_bullet(doc, "Pipeline phải là đồ thị có hướng không chu trình (DAG) và tuân thủ thứ tự category.")
add_bullet(doc, "Config schema là dữ liệu người dùng cấu hình; port schema là hợp đồng artifact giữa các block.")

doc.add_heading("Danh mục artifact", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "Artifact"
table.rows[0].cells[1].text = "Ý nghĩa"
for name, meaning in ARTIFACTS.items():
    cells = table.add_row().cells
    cells[0].text = name
    cells[1].text = meaning
set_table_geometry(table, [2160, 7200])
format_table(table, font_size=9)

doc.add_heading("Mục lục theo nhóm", level=1)
for category in CATEGORY_ORDER:
    number, name, summary = CATEGORY_INFO[category]
    count = sum(1 for b in blocks if b.get("categoryId") == category)
    add_bullet(doc, f"{number}. {name} — {count} block. {summary}")
doc.add_paragraph(f"Tổng cộng: {len(blocks)} block.")

for category_index, category in enumerate(CATEGORY_ORDER):
    number, category_name, summary = CATEGORY_INFO[category]
    if category_index == 0:
        doc.add_page_break()
    doc.add_heading(f"{number}. Nhóm {category_name}", level=1)
    doc.add_paragraph(summary)
    category_blocks = [b for b in blocks if b.get("categoryId") == category]
    for index, block in enumerate(category_blocks, start=1):
        heading = doc.add_heading(f"{number}.{index} {block.get('name')}", level=2)
        heading.paragraph_format.keep_with_next = True
        doc.add_paragraph(block.get("description") or "Chưa có mô tả.")
        add_label_value(doc, "ID", block.get("id"))
        add_label_value(doc, "Code", block.get("code"))
        add_label_value(doc, "Category", block.get("categoryId"))
        add_label_value(doc, "Source", block.get("sourceFile"))
        add_ports_table(doc, (block.get("portSchema") or {}).get("ports", []))
        add_config_table(doc, block.get("configSchema") or {})

doc.add_page_break()
doc.add_heading("Phụ lục A — Quy tắc validation đề xuất", level=1)
for text in (
    "Kiểm tra block ID/code và category tồn tại trong catalog.",
    "Kiểm tra mọi required input đã có edge hợp lệ.",
    "Kiểm tra artifact type hai đầu edge bằng nhau.",
    "Kiểm tra giới hạn multiple và không cho edge trùng.",
    "Kiểm tra required, min, max, minLength và maxLength của config.",
    "Kiểm tra toàn graph không có cycle trước khi Save/Run.",
    "Lưu block version trong graph và migrate config khi definition thay đổi.",
):
    add_bullet(doc, text)

doc.add_heading("Phụ lục B — Ghi chú thiết kế ML", level=1)
add_bullet(doc, "Target column là lựa chọn nghiệp vụ, không nên tự suy luận hoàn toàn từ dữ liệu.")
add_bullet(doc, "Các transformer có bước fit phải chỉ fit trên train set, sau đó áp dụng cùng fitted transformer cho validation/test để tránh data leakage.")
add_bullet(doc, "Cross Validation nên nhận ModelSpec và Folds; runtime tạo estimator mới cho từng fold.")
add_bullet(doc, "Prediction Results xuất Predictions thay vì Metrics để phản ánh đúng dữ liệu đầu ra.")
add_bullet(doc, "Nên bổ sung Regression Metrics và Clustering Metrics khi mở rộng catalog.")

doc.core_properties.title = "Tài liệu hệ thống Block Pipeline ML"
doc.core_properties.subject = "Block definitions, config schema và input/output schema"
doc.core_properties.author = "ML Training Platform"
doc.core_properties.keywords = "ML, pipeline, block, schema, artifact"
doc.save(OUT)
print(f"Created {OUT}")
